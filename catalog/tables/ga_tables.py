import csv
import os
import re
import textwrap
import warnings
from collections import deque
from datetime import date, datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

warnings.filterwarnings("ignore")

COLORS = ["#0072B2", "#E69F00", "#009E73", "#CC79A7", "#56B4E9", "#D55E00", "#F0E442", "#000000"]
SUMMARY_WORDS = {"mean", "average", "avg", "sd", "stdev", "std", "sem", "se", "total", "sum", "median", "count"}
UNIT = re.compile(r"[\(\[].+[\)\]]|%")


def _fail(message):
    return {"ok": False, "error": message}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, (np.floating, float)):
        return None if np.isnan(v) else float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, (pd.Timestamp, datetime, date)):
        return v.isoformat()
    return v


def _fmt(x):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "missing"
    return "%.4g" % x if isinstance(x, (float, np.floating)) else str(x)


def _preview(frame, path, limit=15):
    rows = [[_clean(v) for v in r] for r in frame.head(limit).itertuples(index=False, name=None)]
    return {"columns": [str(c) for c in frame.columns], "rows": rows, "n_rows": int(len(frame)), "path": path}


def _save(frame, stem):
    stem = re.sub(r"\.(csv|tsv|xlsx)$", "", stem, flags=re.I)
    path = os.path.join(os.getcwd(), re.sub(r"[^\w.-]+", "_", stem).strip("_") + ".csv")
    frame.to_csv(path, index=False)
    return path


def _out_dir(sub=""):
    base = os.environ.get("CUVETTE_OUTPUT_DIR") or os.environ.get("CUVETTE_WORK_DIR") or os.getcwd()
    d = os.path.join(base, sub) if sub else base
    os.makedirs(d, exist_ok=True)
    return d


def _unique(d, stem, ext):
    stem = re.sub(r"[^\w.-]+", "-", stem).strip("-") or "table"
    p, n = os.path.join(d, stem + ext), 1
    while os.path.exists(p):
        n += 1
        p = os.path.join(d, "%s-%d%s" % (stem, n, ext))
    return p


def _is_xlsx(path):
    return path.lower().endswith((".xlsx", ".xlsm"))


def text_encoding(path):
    with open(path, "rb") as f:
        raw = f.read()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    for enc in ("utf-8-sig", "cp1252"):
        try:
            raw.decode(enc)
            return enc
        except UnicodeDecodeError:
            pass
    return "latin-1"


def load(path, sheet=None):
    path = os.path.expanduser(path)
    if _is_xlsx(path):
        return pd.read_excel(path, sheet_name=sheet or 0)
    return pd.read_csv(path, sep="\t" if path.lower().endswith((".tsv", ".tab")) else ",", encoding=text_encoding(path))


def _need(frame, cols):
    missing = [c for c in cols if c not in frame.columns]
    if missing:
        return "Unknown column: %s. The columns are: %s." % (", ".join(map(str, missing)), ", ".join(map(str, frame.columns)))
    return None


def _quote_names(expr, columns):
    names = sorted({str(c) for c in columns}, key=len, reverse=True)
    if not names:
        return expr
    pat = re.compile(r"(?<![\w`])(%s)(?![\w`])" % "|".join(re.escape(n) for n in names))
    parts = re.split(r"('[^']*'|\"[^\"]*\"|`[^`]*`)", expr)
    return "".join(p if i % 2 else pat.sub(lambda m: "`%s`" % m.group(1), p) for i, p in enumerate(parts))


def _compute(frame, formulas):
    for name, expr in (formulas or {}).items():
        frame[name] = pd.to_numeric(frame.eval(_quote_names(expr, frame.columns)), errors="coerce")
    return frame


def _excel_formula(expr, frame):
    letters = {str(c): get_column_letter(i + 1) for i, c in enumerate(frame.columns)}
    names = sorted(letters, key=len, reverse=True)
    pat = re.compile(r"`?(%s)`?" % "|".join(re.escape(n) for n in names))
    return "=" + pat.sub(lambda m: letters[m.group(1)] + "2", expr).replace("**", "^")


# ---------------------------------------------------------------- read_table


def _num(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        try:
            return float(s.replace(",", "")) if re.fullmatch(r"[-+]?[\d,]*\.?\d+(e[-+]?\d+)?", s, re.I) else s
        except ValueError:
            return s
    return v


def _grid(path, sheet):
    if _is_xlsx(path):
        wb = load_workbook(path, data_only=True)
        names = wb.sheetnames
        ws = wb[sheet] if sheet else next((wb[n] for n in names if wb[n].max_row > 1 or wb[n]["A1"].value is not None), wb[names[0]])
        cells = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    v = v.strip() or None
                if v is not None:
                    cells[(c.row, c.column)] = v
        merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
        sheets = [{"name": n, "range": wb[n].dimensions, "freeze": wb[n].freeze_panes} for n in names]
        return cells, merges, sheets, ws.title
    sep = "\t" if path.lower().endswith((".tsv", ".tab")) else ","
    raw = pd.read_csv(path, header=None, sep=sep, dtype=str, keep_default_na=False, skip_blank_lines=False, encoding=text_encoding(path))
    cells = {}
    for r, row in enumerate(raw.itertuples(index=False, name=None), 1):
        for c, v in enumerate(row, 1):
            v = _num(v)
            if v is not None:
                cells[(r, c)] = v
    return cells, [], [], None


def _components(cells):
    seen, out = set(), []
    for start in sorted(cells):
        if start in seen:
            continue
        seen.add(start)
        queue, comp = deque([start]), []
        while queue:
            r, c = queue.popleft()
            comp.append((r, c))
            for dr in (-1, 0, 1):
                for dc in (-1, 0, 1):
                    n = (r + dr, c + dc)
                    if n in cells and n not in seen:
                        seen.add(n)
                        queue.append(n)
        rows, cols = [p[0] for p in comp], [p[1] for p in comp]
        out.append({"cells": set(comp), "r0": min(rows), "r1": max(rows), "c0": min(cols), "c1": max(cols)})
    return out


def _addr(r, c):
    return "%s%d" % (get_column_letter(c), r)


def _rng(r0, c0, r1, c1):
    return "%s:%s" % (_addr(r0, c0), _addr(r1, c1))


def _is_text(v):
    return isinstance(v, str)


def _is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _header_row(cells, b):
    all_text = all(_is_text(cells[p]) for p in b["cells"])
    for r in range(b["r0"], min(b["r1"] - 1, b["r0"] + 6) + 1):
        row = [cells[(r, c)] for c in range(b["c0"], b["c1"] + 1) if (r, c) in cells]
        if len(row) < 2 or not all(_is_text(v) for v in row):
            continue
        below = [cells[(r + 1, c)] for c in range(b["c0"], b["c1"] + 1) if (r + 1, c) in cells]
        if all_text or any(not _is_text(v) for v in below):
            return r
    return None


def _spans(label_cells, merges, c0, c1):
    spans = []
    cols = sorted(label_cells)
    for i, c in enumerate(cols):
        merged = next((m for m in merges if m[1] == c and m[0] == label_cells[c][0]), None)
        end = merged[3] if merged else (cols[i + 1] - 1 if i + 1 < len(cols) else c1)
        spans.append((label_cells[c][1], c, min(end, c1)))
    return spans


def _find_tables(cells, merges):
    blocks = _components(cells)
    tables, used = [], set()
    for i, b in enumerate(blocks):
        if b["r1"] - b["r0"] < 1 or b["c1"] - b["c0"] < 1:
            continue
        h = _header_row(cells, b)
        if h is None:
            continue
        used.add(i)
        t = {"block": b, "header": h, "labels": {}}
        for r in range(b["r0"], h):
            for c in range(b["c0"], b["c1"] + 1):
                if (r, c) in cells:
                    t["labels"][c] = (r, str(cells[(r, c)]))
        tables.append(t)
    for t in tables:
        tb = t["block"]
        for i, b in enumerate(blocks):
            if i in used or b["r0"] != b["r1"] or not (t["header"] - 3 <= b["r1"] < t["header"]):
                continue
            if b["c0"] >= tb["c0"] and b["c1"] <= tb["c1"] and all(_is_text(cells[p]) for p in b["cells"]):
                used.add(i)
                for (r, c) in b["cells"]:
                    t["labels"][c] = (r, str(cells[(r, c)]))
    notes = []
    for i, b in enumerate(blocks):
        if i not in used:
            notes.extend({"cell": _addr(r, c), "text": str(cells[(r, c)])} for (r, c) in sorted(b["cells"]))
    return tables, notes


def _build(cells, merges, t):
    b, h = t["block"], t["header"]
    names, cols = [], []
    notes, excluded = [], []
    for c in range(b["c0"], b["c1"] + 1):
        values = [cells.get((r, c)) for r in range(h + 1, b["r1"] + 1)]
        if (h, c) not in cells and sum(v is not None for v in values) <= 2:
            notes.extend({"cell": _addr(r, c), "text": str(cells[(r, c)])} for r in range(h + 1, b["r1"] + 1) if (r, c) in cells)
            continue
        cols.append(c)
        names.append(str(cells.get((h, c), "Column " + get_column_letter(c))))
    rows = []
    for r in range(h + 1, b["r1"] + 1):
        vals = [cells.get((r, c)) for c in cols]
        filled = [v for v in vals if v is not None]
        if not filled:
            continue
        first = next(v for v in vals if v is not None)
        if _is_text(first) and first.strip().lower().rstrip(":") in SUMMARY_WORDS:
            excluded.append({"row": r, "label": first})
            continue
        if len(filled) == 1 and _is_text(first) and len(first.split()) >= 3:
            notes.append({"cell": _addr(r, cols[vals.index(first)]), "text": first})
            continue
        rows.append((r, vals))
    numeric = []
    for j, c in enumerate(cols):
        vals = [v[j] for _, v in rows if v[j] is not None]
        is_num = bool(vals) and sum(_is_number(v) for v in vals) >= 0.6 * len(vals)
        numeric.append(is_num)
        if is_num:
            for r, v in rows:
                if v[j] is not None and not _is_number(v[j]):
                    notes.append({"cell": _addr(r, c), "text": str(v[j])})
                    v[j] = None
    frame = pd.DataFrame([v for _, v in rows], columns=names)
    for j, name in enumerate(names):
        if numeric[j]:
            frame.iloc[:, j] = pd.to_numeric(frame.iloc[:, j], errors="coerce")
    frame = frame.infer_objects()
    labels = {c: v for c, v in t["labels"].items() if c in cols}
    groups = []
    if len(labels) >= 2:
        spans = _spans(labels, merges, b["c0"], b["c1"])
        heads = [[names[cols.index(c)] for c in range(s, e + 1) if c in cols] for _, s, e in spans]
        if all(hd and hd == heads[0] for hd in heads):
            groups = [{"label": lab, "columns": "%s:%s" % (get_column_letter(s), get_column_letter(e)), "cols": [cols.index(c) for c in range(s, e + 1) if c in cols]} for lab, s, e in spans]
    return frame, groups, notes, excluded


def _long(frame, groups):
    grouped = {j for g in groups for j in g["cols"]}
    ids = [frame.columns[j] for j in range(frame.shape[1]) if j not in grouped]
    parts = []
    for g in groups:
        part = frame.iloc[:, g["cols"]].copy()
        part.columns = [str(c) for c in part.columns]
        part = part.dropna(how="all")
        part.insert(0, "Group", g["label"])
        for k, idc in enumerate(ids):
            part.insert(k, idc, frame.loc[part.index, idc])
        parts.append(part)
    return pd.concat(parts, ignore_index=True)


def _col_types(frame):
    out = []
    for j, c in enumerate(frame.columns):
        s = frame.iloc[:, j]
        kind = "number" if pd.api.types.is_numeric_dtype(s) else "date" if pd.api.types.is_datetime64_any_dtype(s) else "text"
        out.append({"name": str(c), "type": kind, "n_missing": int(s.isna().sum())})
    return out


def read_table(path, sheet=None, table=None, range=None, groups_to_rows=True):
    path = os.path.expanduser(path)
    if not os.path.exists(path):
        return _fail("No file at %s." % path)
    cells, merges, sheets, sheet_name = _grid(path, sheet)
    if not cells:
        return _fail("The sheet is empty.")
    tables, notes = _find_tables(cells, merges)
    if range:
        r0, c0, r1, c1 = _parse_range(range)
        sub = {p: v for p, v in cells.items() if r0 <= p[0] <= r1 and c0 <= p[1] <= c1}
        tables = [{"block": {"cells": set(sub), "r0": r0, "r1": r1, "c0": c0, "c1": c1}, "header": r0, "labels": {}}]
    if not tables:
        cells_text = "; ".join('%s "%s"' % (n["cell"], n["text"][:80]) for n in notes[:20])
        return _fail("No table with a header row and two or more columns found. The cells are: %s. Give range=C5:G20 to read a range." % cells_text)
    found = []
    for k, t in enumerate(tables, 1):
        frame, groups, tnotes, excluded = _build(cells, merges, t)
        b = t["block"]
        found.append({
            "index": k, "range": _rng(t["header"], b["c0"], b["r1"], b["c1"]),
            "header_row": t["header"], "n_rows": int(len(frame)), "columns": _col_types(frame),
            "label_row": next(iter(t["labels"].values()))[0] if t["labels"] else None,
            "labels": [v[1] for _, v in sorted(t["labels"].items())],
            "groups": [{"label": g["label"], "columns": g["columns"]} for g in groups],
            "excluded_rows": excluded, "_frame": frame, "_groups": groups,
        })
        notes.extend(tnotes)
    pick = int(table) if table else max(found, key=lambda f: f["_frame"].size)["index"]
    if not 1 <= pick <= len(found):
        return _fail("There are %d tables. Give table as a number from 1 to %d." % (len(found), len(found)))
    chosen = found[pick - 1]
    frame = chosen["_frame"]
    long = bool(chosen["_groups"]) and groups_to_rows
    if long:
        frame = _long(frame, chosen["_groups"])
    elif chosen["_groups"]:
        names = [str(c) for c in frame.columns]
        for gr in chosen["_groups"]:
            for j in gr["cols"]:
                names[j] = "%s %s" % (gr["label"], names[j])
        frame.columns = names
    stem = os.path.splitext(os.path.basename(path))[0] + ("-" + sheet_name if sheet_name else "") + "-table%d" % pick
    out = _save(frame, stem)
    for f in found:
        f.pop("_frame")
        f.pop("_groups")
    notes = sorted(notes, key=lambda n: (int(re.sub(r"\D", "", n["cell"])), n["cell"]))
    lines = []
    if sheets:
        lines.append("Sheets: %s. Read sheet %s." % (", ".join("%s (%s)" % (s["name"], s["range"]) for s in sheets), sheet_name))
    for f in found:
        desc = "Table %d at %s: header row %d, %d rows, columns %s." % (
            f["index"], f["range"], f["header_row"], f["n_rows"], ", ".join("%s (%s)" % (c["name"], c["type"]) for c in f["columns"]))
        if f["groups"]:
            desc += " Group labels in row %d: %s." % (f["label_row"], ", ".join("%s over %s" % (g["label"], g["columns"]) for g in f["groups"]))
        elif f["labels"]:
            desc += " Title above it: %s." % "; ".join(f["labels"])
        if f["excluded_rows"]:
            desc += " Left out summary rows: %s." % ", ".join("%d (%s)" % (e["row"], e["label"]) for e in f["excluded_rows"])
        lines.append(desc)
    if notes:
        lines.append("Cells that look like notes (not data): %s." % "; ".join('%s "%s"' % (n["cell"], n["text"][:60]) for n in notes))
    lines.append("Saved table %d%s as %s (%d rows, columns %s)." % (
        pick, ", with one row per sample and group" if long else "", out, len(frame), ", ".join(map(str, frame.columns))))
    if len(found) > 1:
        lines.append("Give table=N to read another table.")
    rng = chosen["range"]
    route = "In Excel: open %s%s. Select %s (the header is row %d). Copy it to a new sheet with Home > Paste > Values." % (
        os.path.basename(path), ", sheet %s" % sheet_name if sheet_name else "", rng, chosen["header_row"])
    if long:
        route += " To get one row per sample and group: copy each group block (%s) under the other, and add a column Group with the label of the block." % ", ".join(g["columns"] for g in chosen["groups"])
    if chosen["excluded_rows"]:
        route += " Do not copy the summary rows %s." % ", ".join(str(e["row"]) for e in chosen["excluded_rows"])
    return {
        "ok": True,
        "summary": " ".join(lines),
        "metrics": {"n_tables": len(found), "n_notes": len(notes), "n_rows": int(len(frame)), "n_columns": int(frame.shape[1]), "n_sheets": len(sheets)},
        "path": out, "sheet": sheet_name, "sheets": sheets, "tables": found, "notes": notes, "columns": _col_types(frame),
        "table": _preview(frame, out),
        "outputs": [{"path": out, "kind": "table", "name": os.path.basename(out)}],
        "manualRoute": route,
        "fidelity": "close",
    }


def _parse_range(text):
    m = re.fullmatch(r"\s*([A-Za-z]+)(\d+)\s*:\s*([A-Za-z]+)(\d+)\s*", text)
    if not m:
        raise ValueError("A range looks like C5:G20.")
    return int(m.group(2)), column_index_from_string(m.group(1).upper()), int(m.group(4)), column_index_from_string(m.group(3).upper())


# ------------------------------------------------------------ summarize_table


def summarize_table(path, values=None, group_by=None, ratios=None, sheet=None):
    frame = load(path, sheet)
    group_by = [group_by] if isinstance(group_by, str) else list(group_by or [])
    values = [values] if isinstance(values, str) else list(values or [])
    try:
        frame = _compute(frame, ratios)
    except Exception as e:
        return _fail("A ratio formula failed: %s. Write it like: Red dots / Total cells * 100." % e)
    problem = _need(frame, group_by + values)
    if problem:
        return _fail(problem)
    if not values:
        values = [c for c in frame.columns if c not in group_by and pd.api.types.is_numeric_dtype(frame[c]) and c not in (ratios or {})]
    values = values + [r for r in (ratios or {}) if r not in values]
    rows = []
    groups = frame.groupby(group_by, sort=False, dropna=False) if group_by else [((), frame)]
    for key, part in groups:
        key = key if isinstance(key, tuple) else (key,)
        for v in values:
            s = pd.to_numeric(part[v], errors="coerce").dropna()
            n = int(len(s))
            sd = float(s.std(ddof=1)) if n > 1 else float("nan")
            q1, q3 = (float(s.quantile(0.25)), float(s.quantile(0.75))) if n else (float("nan"),) * 2
            rows.append(list(key) + [v, n, float(s.mean()) if n else float("nan"), sd, sd / np.sqrt(n) if n > 1 else float("nan"),
                                     float(s.median()) if n else float("nan"), q1, q3, q3 - q1,
                                     float(s.min()) if n else float("nan"), float(s.max()) if n else float("nan")])
    out = pd.DataFrame(rows, columns=group_by + ["variable", "n", "mean", "sd", "sem", "median", "q1", "q3", "iqr", "min", "max"])
    saved = _save(out, "summary")
    parts = []
    for r in rows[:8]:
        label = " ".join(map(str, r[:len(group_by)]))
        parts.append("%s%s: mean %s (SD %s, SEM %s, n %d)" % (label + " " if label else "", r[len(group_by)], _fmt(r[len(group_by) + 2]),
                                                           _fmt(r[len(group_by) + 3]), _fmt(r[len(group_by) + 4]), r[len(group_by) + 1]))
    by = " by " + ", ".join(group_by) if group_by else ""
    summary = "Summarized %s%s over %d rows. %s%s." % (", ".join(values), by, len(frame), "; ".join(parts), "; ..." if len(rows) > 8 else "")
    if ratios:
        rows_path = _save(frame, "rows-with-ratios")
        summary += " Each ratio is computed per row, then summarized. The rows with the ratio columns are in %s; use it for make_plot." % rows_path
    route = []
    for name, expr in (ratios or {}).items():
        route.append('add a column "%s" with %s and fill down' % (name, _excel_formula(expr, frame)))
    route.append("Insert > PivotTable; Rows: %s; Values: each column, with Value Field Settings > Average, Count and StdDev" % (", ".join(group_by) or "none"))
    route.append("SEM = STDEV.S(range)/SQRT(COUNT(range)); median = MEDIAN(range); quartiles = QUARTILE.INC(range,1) and QUARTILE.INC(range,3); IQR = Q3 - Q1")
    return {
        "ok": True,
        "summary": summary,
        "metrics": {"n_rows": int(len(frame)), "n_groups": len(rows) // max(len(values), 1), "n_values": len(values)},
        "path": saved,
        "table": _preview(out, saved, 30),
        "outputs": [{"path": saved, "kind": "table", "name": os.path.basename(saved)}],
        "manualRoute": "In Excel: " + "; ".join(route) + ".",
        "fidelity": "exact",
    }


# -------------------------------------------------------------- reshape_table


def reshape_table(path, sheet=None, join=None, compute=None, filter=None, pivot=None, melt=None, select=None, rename=None, sort=None, name=None):
    frame = load(path, sheet)
    n_in = len(frame)
    steps, route, metrics = [], [], {}
    try:
        if join:
            other = load(join["path"], join.get("sheet"))
            on = join["on"]
            keys = [on] if isinstance(on, str) else list(on)
            problem = _need(frame, keys) or _need(other, keys)
            if problem:
                return _fail("Join: " + problem)
            how = join.get("how", "left")
            merged = frame.merge(other, on=keys, how=how, indicator=True, suffixes=("", " (2)"))
            unmatched = int((merged["_merge"] == "left_only").sum())
            frame = merged.drop(columns="_merge")
            metrics["n_unmatched"] = unmatched
            steps.append("joined %s on %s (%s join; %d rows had no match%s)" % (
                os.path.basename(join["path"]), ", ".join(keys), how, unmatched,
                "; the second table has repeated keys, so some rows doubled" if how == "left" and len(frame) > n_in else ""))
            route.append("add the columns of the second table with =XLOOKUP(key cell, second table key column, second table value column, \"\") and fill down (or Data > Get Data > Combine Queries > Merge, %s outer)" % how.capitalize())
        if compute:
            frame = _compute(frame, compute)
            steps.append("added %s" % ", ".join('"%s" = %s' % kv for kv in compute.items()))
            route.extend('add a column "%s" with %s and fill down' % (k, _excel_formula(v, frame)) for k, v in compute.items())
        if filter:
            before = len(frame)
            frame = frame.query(_quote_names(filter, frame.columns)).reset_index(drop=True)
            steps.append("kept %d of %d rows where %s" % (len(frame), before, filter))
            route.append("Data > Filter, set the filter to keep rows where %s, then copy the visible rows to a new sheet" % filter)
        if pivot:
            idx = pivot["index"]
            idx = [idx] if isinstance(idx, str) else list(idx)
            vals = pivot["values"]
            dup = int(frame.duplicated(idx + [pivot["columns"]]).sum())
            if dup:
                return _fail("Pivot: %d rows repeat the same %s and %s. Filter or summarize the table first." % (dup, ", ".join(idx), pivot["columns"]))
            wide = frame.pivot(index=idx, columns=pivot["columns"], values=vals)
            if isinstance(wide.columns, pd.MultiIndex):
                wide.columns = [" ".join(map(str, c)) for c in wide.columns]
            wide.columns = [str(c) for c in wide.columns]
            frame = wide.reset_index()
            steps.append("pivoted to wide: one row per %s, one column per %s" % (", ".join(idx), pivot["columns"]))
            route.append("Insert > PivotTable; Rows: %s; Columns: %s; Values: %s (Sum of one value is the value)" % (", ".join(idx), pivot["columns"], vals if isinstance(vals, str) else ", ".join(vals)))
        if melt:
            ids = melt.get("id_columns") or []
            frame = frame.melt(id_vars=ids, value_vars=melt.get("value_columns"), var_name=melt.get("name", "variable"), value_name=melt.get("value", "value"))
            steps.append("made long: one row per %s and column" % ", ".join(ids))
            route.append("select the table; Data > From Table/Range; select %s; Transform > Unpivot Other Columns; Home > Close & Load" % ", ".join(ids))
        if select:
            problem = _need(frame, select)
            if problem:
                return _fail(problem)
            frame = frame[list(select)]
            steps.append("kept columns %s" % ", ".join(select))
            route.append("delete the other columns")
        if rename:
            frame = frame.rename(columns=rename)
            steps.append("renamed %s" % ", ".join("%s to %s" % kv for kv in rename.items()))
            route.append("type the new header names")
        if sort:
            keys = [s.lstrip("-") for s in sort]
            frame = frame.sort_values(keys, ascending=[not s.startswith("-") for s in sort], kind="stable").reset_index(drop=True)
            steps.append("sorted by %s" % ", ".join(sort))
            route.append("Data > Sort, by %s" % ", ".join("%s (%s)" % (k, "Largest to Smallest" if s.startswith("-") else "Smallest to Largest") for k, s in zip(keys, sort)))
    except KeyError as e:
        return _fail("Unknown column or key: %s. The columns are: %s." % (e, ", ".join(map(str, frame.columns))))
    except Exception as e:
        return _fail("%s: %s" % (type(e).__name__, e))
    if not steps:
        return _fail("Give at least one step: join, compute, filter, pivot, melt, select, rename or sort.")
    saved = _save(frame, name or os.path.splitext(os.path.basename(path))[0] + "-reshaped")
    metrics.update({"n_rows": int(len(frame)), "n_columns": int(frame.shape[1]), "n_rows_in": n_in})
    return {
        "ok": True,
        "summary": "From %d rows: %s. The new table has %d rows and columns %s. Saved as %s." % (n_in, "; ".join(steps), len(frame), ", ".join(map(str, frame.columns)), saved),
        "metrics": metrics,
        "path": saved, "columns": _col_types(frame),
        "table": _preview(frame, saved),
        "outputs": [{"path": saved, "kind": "table", "name": os.path.basename(saved)}],
        "manualRoute": "In Excel: " + "; then ".join(route) + ".",
        "fidelity": "close",
    }


# ------------------------------------------------------------------ make_plot


def _style():
    plt.rcParams.update({
        "font.size": 10, "axes.labelsize": 10, "xtick.labelsize": 9, "ytick.labelsize": 9, "legend.fontsize": 9,
        "font.family": "sans-serif", "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
        "axes.spines.top": False, "axes.spines.right": False, "axes.linewidth": 0.8,
        "svg.fonttype": "none", "savefig.dpi": 300, "legend.frameon": False,
    })


def _err(s, kind):
    n = s.count()
    if n < 2:
        return 0.0
    sd = s.std(ddof=1)
    return float(sd / np.sqrt(n) if kind == "SEM" else sd)


def make_plot(path, kind, y, y_label, x=None, x_label=None, group=None, error_bars="SD", points=True, fit=False, bins=None, title=None, file_name=None, sheet=None):
    frame = load(path, sheet)
    kind = kind.lower()
    if kind not in ("bar", "box", "scatter", "line", "histogram"):
        return _fail("kind must be bar, box, scatter, line or histogram.")
    if kind in ("bar", "box") and not x:
        return _fail("A %s plot needs x, the group column." % kind)
    if kind in ("scatter", "line") and not x:
        return _fail("A %s plot needs x." % kind)
    problem = _need(frame, [c for c in (x, y, group) if c])
    if problem:
        return _fail(problem)
    need_unit = [("y_label", y_label)] + ([("x_label", x_label)] if kind in ("scatter", "line") else [])
    for arg, label in need_unit:
        if not label or not UNIT.search(label):
            return _fail("%s needs the unit in parentheses, for example 'Red dots per cell (%%)'. For a count, write '(count)'." % arg)
    error_bars = (error_bars or "SD").upper()
    frame[y] = pd.to_numeric(frame[y], errors="coerce")
    frame = frame.dropna(subset=[y])
    _style()
    rng = np.random.default_rng(0)
    if kind in ("bar", "box"):
        cats = list(dict.fromkeys(frame[x].astype(str)))
        width = max(2.2, 0.9 * len(cats) + 1.0)
    else:
        width = 3.5
    fig, ax = plt.subplots(figsize=(width, 3.0), layout="constrained")
    metrics, desc = {"n": int(len(frame))}, ""
    shows_err = kind == "bar"
    if kind in ("bar", "box"):
        data = [frame.loc[frame[x].astype(str) == c, y] for c in cats]
        pos = np.arange(len(cats))
        colors = [COLORS[i % len(COLORS)] for i in range(len(cats))]
        if kind == "bar":
            means = [d.mean() for d in data]
            errs = [_err(d, error_bars) for d in data]
            ax.bar(pos, means, width=0.6, color=colors, alpha=0.35 if points else 0.85, edgecolor=colors, linewidth=1.0, zorder=1)
            ax.errorbar(pos, means, yerr=errs, fmt="none", ecolor="black", elinewidth=1.0, capsize=4, zorder=3)
            desc = "Bar plot of the mean %s by %s with %s error bars%s: %s" % (
                y, x, error_bars, " and individual points" if points else "",
                ", ".join("%s %s ± %s (n = %d)" % (c, _fmt(m), _fmt(e), len(d)) for c, m, e, d in zip(cats, means, errs, data)))
            for c, m in zip(cats, means):
                metrics["mean_" + c] = float(m)
        else:
            ax.boxplot(data, positions=pos, widths=0.55, showfliers=not points, medianprops={"color": "black"},
                       patch_artist=True, boxprops={"facecolor": "white"})
            for patch, col in zip(ax.patches, colors):
                patch.set_edgecolor(col)
            desc = "Box plot of %s by %s (median, quartiles, whiskers to 1.5 IQR)%s: medians %s" % (
                y, x, " with individual points" if points else "",
                ", ".join("%s %s (n = %d)" % (c, _fmt(d.median()), len(d)) for c, d in zip(cats, data)))
            for c, d in zip(cats, data):
                metrics["median_" + c] = float(d.median())
        if points:
            for i, d in enumerate(data):
                ax.scatter(i + rng.uniform(-0.12, 0.12, len(d)), d, s=16, color=colors[i], edgecolor="black", linewidth=0.4, zorder=4)
        ax.set_xticks(pos, ["%s\nn = %d" % (c, len(d)) for c, d in zip(cats, data)])
        ax.set_xlim(-0.6, len(cats) - 0.4)
        if kind == "bar" and min(d.min() for d in data if len(d)) >= 0:
            ax.set_ylim(bottom=0)
        metrics["n_groups"] = len(cats)
    elif kind == "histogram":
        sets = [(str(g), part[y]) for g, part in frame.groupby(group, sort=False)] if group else [(None, frame[y])]
        edges = np.histogram_bin_edges(frame[y], bins=int(bins) if bins else "auto")
        for i, (g, s) in enumerate(sets):
            ax.hist(s, bins=edges, color=COLORS[i % len(COLORS)], alpha=0.5 if group else 0.85, edgecolor="white", linewidth=0.5,
                    label="%s (n = %d)" % (g, len(s)) if group else None)
        ax.set_ylabel("Count")
        ax.set_xlabel(y_label)
        if not group:
            ax.text(0.98, 0.95, "n = %d" % len(frame), transform=ax.transAxes, ha="right", va="top")
        metrics["n_bins"] = len(edges) - 1
        metrics["bin_width"] = float(edges[1] - edges[0])
        desc = "Histogram of %s (n = %d) in %d bins of width %s; median %s" % (y, len(frame), len(edges) - 1, _fmt(edges[1] - edges[0]), _fmt(frame[y].median()))
    else:
        frame[x] = pd.to_numeric(frame[x], errors="coerce")
        frame = frame.dropna(subset=[x])
        metrics["n"] = int(len(frame))
        sets = [(str(g), part) for g, part in frame.groupby(group, sort=False)] if group else [(None, frame)]
        parts = []
        for i, (g, part) in enumerate(sets):
            col = COLORS[i % len(COLORS)]
            lab = "%s (n = %d)" % (g, len(part)) if g is not None else None
            if kind == "scatter":
                ax.scatter(part[x], part[y], s=18, color=col, edgecolor="black", linewidth=0.4, label=lab, zorder=3)
                if fit and len(part) > 2:
                    slope, icpt = np.polyfit(part[x], part[y], 1)
                    r = float(np.corrcoef(part[x], part[y])[0, 1])
                    xs = np.linspace(part[x].min(), part[x].max(), 50)
                    ax.plot(xs, slope * xs + icpt, color=col, linewidth=1.2)
                    key = "" if g is None else "_" + g
                    metrics["slope" + key], metrics["intercept" + key], metrics["r" + key] = float(slope), float(icpt), r
                    parts.append("%sslope %s, r = %s, n = %d" % ("%s: " % g if g else "", _fmt(slope), "%.3f" % r, len(part)))
                else:
                    parts.append("%sn = %d" % ("%s " % g if g else "", len(part)))
            else:
                agg = part.groupby(x)[y].agg(["mean", "count", "std"]).reset_index()
                errs = agg["std"].fillna(0) / (np.sqrt(agg["count"]) if error_bars == "SEM" else 1)
                reps = bool((agg["count"] > 1).any())
                shows_err = shows_err or reps
                ax.errorbar(agg[x], agg["mean"], yerr=errs if reps else None, color=col, marker="o", markersize=4, linewidth=1.2, capsize=3, label=lab)
                if points and reps:
                    ax.scatter(part[x], part[y], s=10, color=col, alpha=0.5, zorder=2)
                parts.append("%s%d points" % ("%s " % g if g else "", len(agg)))
        if group:
            ax.legend()
        elif kind == "scatter":
            ax.text(0.02, 0.95, "n = %d" % len(frame) + ("\nr = %.3f" % metrics["r"] if "r" in metrics else ""), transform=ax.transAxes, va="top")
        ax.set_xlabel(x_label)
        if kind == "scatter":
            desc = "Scatter plot of %s against %s%s: %s" % (y, x, " with a linear fit" if fit else "", "; ".join(parts))
        else:
            desc = "Line plot of the mean %s against %s%s: %s" % (y, x, " with %s error bars" % error_bars if shows_err else "", "; ".join(parts))
    if kind != "histogram":
        ax.set_ylabel(y_label)
        if shows_err and not title:
            ax.text(1.0, 1.0, "mean ± %s" % error_bars, transform=ax.transAxes, ha="right", va="bottom", fontsize=8)
        if x_label and kind in ("bar", "box"):
            ax.set_xlabel(x_label)
    if title:
        ax.set_title(textwrap.fill(title + (" (mean ± %s)" % error_bars if shows_err else ""), int(width * 11)), fontsize=10)
    d = _out_dir("figures")
    stem = file_name or "%s-%s" % (kind, y)
    png = _unique(d, os.path.splitext(os.path.basename(stem))[0], ".png")
    svg = png[:-4] + ".svg"
    fig.savefig(png, dpi=300)
    if not os.path.exists(svg) or os.path.getmtime(svg) < os.path.getmtime(png):
        fig.savefig(svg)
    plt.close(fig)
    w, h = plt.imread(png).shape[1::-1]
    metrics.update({"width_px": int(w), "height_px": int(h), "png_bytes": os.path.getsize(png), "svg_bytes": os.path.getsize(svg)})
    return {
        "ok": True,
        "summary": desc + ". Saved %s and %s." % (png, os.path.basename(svg)),
        "metrics": metrics,
        "png": png, "svg": svg,
        "outputs": [{"path": png, "kind": "image", "name": os.path.basename(png)}, {"path": svg, "kind": "file", "name": os.path.basename(svg)}],
        "manualRoute": _plot_route(kind, x, y, error_bars, points, fit, metrics),
        "fidelity": "close",
    }


def _plot_route(kind, x, y, error_bars, points, fit, metrics):
    err = "=STDEV.S(range)" if error_bars == "SD" else "=STDEV.S(range)/SQRT(COUNT(range))"
    custom = ("add error bars: Chart Design > Add Chart Element > Error Bars > More Error Bar Options > Custom > Specify Value, "
              "and pick the %s cells for both + and - (the built-in Standard Deviation option uses one SD across all bars, not each group's SD)" % error_bars)
    dots = ("; add the points: put each group's values in a column next to an x column of 1, 2, 3... (plus small offsets), "
            "Chart Design > Select Data > Add, then Change Series Chart Type > Scatter")
    if kind == "bar":
        return ("In Excel: make a small table with one row per %s: =AVERAGE(range) and %s (%s). Select the means; Insert > Chart > Clustered Column; %s%s. "
                "Axis titles: Chart Design > Add Chart Element > Axis Titles. Delete the gridlines." % (x, err, error_bars, custom, dots if points else ""))
    if kind == "box":
        return ("In Excel: put each %s's values in its own column; Insert > Chart > Box and Whisker; Format Data Series > Quartile Calculation > "
                "Inclusive median (this plot uses the inclusive method); untick Show outlier points%s." % (x, " and tick Show inner points" if points else ""))
    if kind == "scatter":
        return ("In Excel: select the %s and %s columns; Insert > Chart > Scatter%s." % (
            x, y, "; Chart Design > Add Chart Element > Trendline > Linear; Format Trendline > Display Equation and Display R-squared; r = CORREL(x range, y range)" if fit else ""))
    if kind == "line":
        return "In Excel: make a table of the mean %s for each %s (PivotTable, Average); Insert > Chart > Scatter with Straight Lines and Markers; %s." % (y, x, custom)
    return "In Excel: select the %s column; Insert > Chart > Histogram; Format Axis > Bin width %s." % (y, _fmt(metrics.get("bin_width", 0)))


# ----------------------------------------------------------------- read_plate

OVERFLOW = re.compile(r"^(ovr|over|sat|\*+$|>|\++$)", re.I)
NUMBER = re.compile(r"[-+]?(\d+\.?\d*|\.\d+)(e[-+]?\d+)?", re.I)
WELL_COLUMNS = {"well", "well id", "wellid", "well position", "position"}


def _raw(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return None
    if isinstance(v, str):
        return v.strip() or None
    if isinstance(v, (bool, int, float)):
        return v
    return str(v)


def _number(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and NUMBER.fullmatch(v):
        return float(v)
    return None


def _int_of(v):
    n = _number(v)
    return int(n) if n is not None and n.is_integer() else None


def _text_of(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _plate_sheets(path, sheet=None):
    if _is_xlsx(path):
        wb = load_workbook(path, data_only=True)
        if sheet and sheet not in wb.sheetnames:
            raise ValueError("No sheet %s. The sheets are: %s." % (sheet, ", ".join(wb.sheetnames)))
        out = []
        for n in [sheet] if sheet else wb.sheetnames:
            cells = {}
            for row in wb[n].iter_rows():
                for c in row:
                    v = _raw(c.value)
                    if v is not None:
                        cells[(c.row, c.column)] = v
            out.append((n, cells))
        return out
    if path.lower().endswith(".xls"):
        try:
            book = pd.read_excel(path, sheet_name=sheet if sheet else None, header=None, dtype=object)
        except ImportError:
            raise ValueError("Python cannot read .xls files here (the xlrd package is missing). Save the file as .xlsx in Excel.")
        book = {sheet: book} if sheet else book
        out = []
        for n, f in book.items():
            cells = {}
            for r, row in enumerate(f.itertuples(index=False, name=None), 1):
                for c, v in enumerate(row, 1):
                    v = _raw(v)
                    if v is not None:
                        cells[(r, c)] = v
            out.append((n, cells))
        return out
    sep = "\t" if path.lower().endswith((".tsv", ".tab")) else ","
    cells = {}
    with open(path, encoding=text_encoding(path), newline="") as f:
        for r, row in enumerate(csv.reader(f, delimiter=sep), 1):
            for c, v in enumerate(row, 1):
                v = _raw(v)
                if v is not None:
                    cells[(r, c)] = v
    return [(os.path.basename(path), cells)]


def _plate_blocks(cells):
    found = []
    for (r, c) in sorted(cells):
        if _int_of(cells[(r, c)]) != 1 or any(b["head"] <= r <= b["r1"] and b["c0"] <= c <= b["c1"] for b in found):
            continue
        n = 1
        while _int_of(cells.get((r, c + n))) == n + 1:
            n += 1
        if n < 3:
            continue
        for lc in (c - 1, c - 2):
            if lc < 1 or str(cells.get((r + 1, lc), "")).upper() != "A":
                continue
            k = 0
            while str(cells.get((r + 1 + k, lc), "")).upper() == get_column_letter(k + 1):
                k += 1
            if k >= 2:
                found.append({"head": r, "lc": lc, "c0": c, "c1": c + n - 1, "r0": r + 1, "r1": r + k, "by_order": False})
                break
        else:
            m = 0
            while any(_number(cells.get((r + 1 + m, cc))) is not None for cc in range(c, c + n)):
                m += 1
            k = 16 if m >= 16 else 8 if m >= 8 else 0
            if k and n in (12, 24):
                found.append({"head": r, "lc": max(c - 1, 1), "c0": c, "c1": c + n - 1, "r0": r + 1, "r1": r + k, "by_order": True})
    return found


def _plate_label(cells, blocks, b):
    side = [o for o in blocks if o is not b and o["head"] == b["head"]]
    floor = max([o["r1"] for o in blocks if o["r1"] < b["head"]], default=0)
    lo = max([o["c1"] + 1 for o in side if o["c1"] < b["c0"]], default=1)
    hi = min([o["lc"] - 1 for o in side if o["c0"] > b["c1"]], default=10**6)
    lines, r, skipped_blank = [], b["head"] - 1, False
    while r > floor and len(lines) < 2:
        parts = [_text_of(cells[(r, c)]) for c in sorted({c for (rr, c) in cells if rr == r}) if lo <= c <= hi]
        if parts:
            lines.insert(0, " ".join(parts))
        elif lines or skipped_blank:
            break
        else:
            skipped_blank = True
        r -= 1
    return " | ".join(lines)


def _well_key(text):
    m = re.fullmatch(r"\s*([A-Za-z]{1,2})\s*0*(\d{1,2})\s*", str(text))
    return m.group(1).upper() + m.group(2) if m else str(text).strip()


def _read_map(map_path, map_sheet, well_column, map_name):
    for name, cells in _plate_sheets(map_path, map_sheet):
        blocks = _plate_blocks(cells)
        if not blocks:
            continue
        b = blocks[0]
        entries = {}
        for i in range(b["r1"] - b["r0"] + 1):
            for j in range(b["c1"] - b["c0"] + 1):
                v = cells.get((b["r0"] + i, b["c0"] + j))
                if v is not None:
                    entries[get_column_letter(i + 1) + str(j + 1)] = _text_of(v)
        frame = pd.DataFrame({"_key": list(entries), map_name: list(entries.values())})
        extra = "; the sheet has %d blocks and the tool used the first" % len(blocks) if len(blocks) > 1 else ""
        return frame, [map_name], "grid on %s at %s with text in %d wells%s" % (name, _rng(b["head"], b["lc"], b["r1"], b["c1"]), len(frame), extra)
    if _is_xlsx(map_path) or map_path.lower().endswith(".xls"):
        frame = pd.read_excel(map_path, sheet_name=map_sheet or 0, dtype=str, keep_default_na=False)
    else:
        frame = pd.read_csv(map_path, sep="\t" if map_path.lower().endswith((".tsv", ".tab")) else ",", dtype=str, keep_default_na=False, encoding=text_encoding(map_path))
    frame.columns = [str(c).strip() for c in frame.columns]
    col = well_column or next((c for c in frame.columns if c.lower() in WELL_COLUMNS), None)
    if col not in frame.columns:
        raise ValueError("The plate map has no row-letter and column-number grid and no well column. Its columns are: %s. Give map_well_column." % ", ".join(frame.columns))
    frame["_key"] = frame[col].map(_well_key)
    frame = frame[frame["_key"] != ""]
    dup = frame["_key"][frame["_key"].duplicated()].unique()
    if len(dup):
        raise ValueError("The plate map lists these wells more than once: %s." % ", ".join(dup[:10]))
    cols = [c for c in frame.columns if c not in (col, "_key")]
    return frame[["_key"] + cols], cols, "long table with the well column %s and %d wells" % (col, len(frame))


def _list_wells(wells, limit=10):
    wells = list(wells)
    return ", ".join(wells[:limit]) + (" and %d more" % (len(wells) - limit) if len(wells) > limit else "")


def _well_order(w):
    m = re.fullmatch(r"([A-Z]+)(\d+)", w)
    return (len(m.group(1)), m.group(1), int(m.group(2))) if m else (99, w, 0)


def read_plate(path, sheet=None, plate_map=None, map_sheet=None, map_well_column=None, map_name="map"):
    path = os.path.expanduser(path)
    if not os.path.exists(path):
        return _fail("No file at %s." % path)
    map_path = os.path.expanduser(plate_map) if plate_map else (path if map_sheet else None)
    try:
        sheets = _plate_sheets(path, sheet)
        mapping = _read_map(map_path, map_sheet, map_well_column, map_name) if map_path else None
    except ValueError as e:
        return _fail("Cannot read the plate file or the plate map: %s" % e)
    if mapping and os.path.abspath(map_path) == os.path.abspath(path) and map_sheet and not sheet:
        sheets = [s for s in sheets if s[0] != map_sheet]
    rows, blocks, skipped = [], [], []
    map_keys = set(mapping[0]["_key"]) if mapping else set()
    for plate, cells in sheets:
        found = _plate_blocks(cells)
        used = set()
        for k, b in enumerate(found, 1):
            name = _plate_label(cells, found, b) or "Block %d" % k
            if name in used:
                name = "%s (%d)" % (name, k)
            used.add(name)
            rng = _rng(b["head"], b["lc"], b["r1"], b["c1"])
            grid = [[cells.get((b["r0"] + i, b["c0"] + j)) for j in range(b["c1"] - b["c0"] + 1)] for i in range(b["r1"] - b["r0"] + 1)]
            if not any(_number(v) is not None for line in grid for v in line):
                skipped.append("%s on %s" % (rng, plate))
                continue
            info = {"plate": plate, "block": name, "range": rng, "n_rows": len(grid), "n_cols": len(grid[0]), "n_wells": 0, "n_values": 0, "n_empty": 0, "n_text": 0, "n_overflow": 0, "letters_from_order": b["by_order"], "overflow_wells": [], "text_wells": [], "empty_wells": []}
            have = set()
            for i, line in enumerate(grid):
                for j, v in enumerate(line):
                    well = get_column_letter(i + 1) + str(j + 1)
                    x = _number(v)
                    flag = ""
                    info["n_wells"] += 1
                    if x is not None:
                        info["n_values"] += 1
                        have.add(well)
                    elif v is None:
                        flag = "empty"
                        info["n_empty"] += 1
                        info["empty_wells"].append(well)
                    elif OVERFLOW.match(str(v)):
                        flag = str(v)
                        info["n_overflow"] += 1
                        info["overflow_wells"].append('%s "%s"' % (well, v))
                    else:
                        flag = str(v)
                        info["n_text"] += 1
                        info["text_wells"].append('%s "%s"' % (well, v))
                    rows.append({"plate": plate, "block": name, "well": well, "row": get_column_letter(i + 1), "col": j + 1, "value": x, "flag": flag})
            info["map_without_value"] = sorted(map_keys - have, key=_well_order)
            info["n_map_without_value"] = len(info["map_without_value"])
            blocks.append(info)
    if not blocks:
        extra = " Text-only grids that the tool skipped: %s." % "; ".join(skipped) if skipped else ""
        return _fail("No plate block found. A block has a row with the numbers 1, 2, 3 and so on, and the row letters A, B, C and so on in the column to its left.%s" % extra)
    frame = pd.DataFrame(rows)
    n_map_cols = 0
    if mapping:
        mframe, mcols, _ = mapping
        taken = set(frame.columns) | {"_key"}
        names = {c: ("map_" + c if c in taken else c) for c in mcols}
        frame = frame.merge(mframe.rename(columns=names), how="left", left_on="well", right_on="_key").drop(columns="_key")
        n_map_cols = len(mcols)
    out = _save(frame, os.path.splitext(os.path.basename(path))[0] + "-plate")
    lines = ["Sheets read: %s." %", ".join(s[0] for s in sheets), "Found %d plate blocks." % len(blocks)]
    for b in blocks:
        miss = []
        if b["n_empty"]:
            miss.append("%d empty (%s)" % (b["n_empty"], _list_wells(b["empty_wells"])))
        if b["n_overflow"]:
            miss.append("%d overflow (%s)" % (b["n_overflow"], _list_wells(b["overflow_wells"])))
        if b["n_text"]:
            miss.append("%d other text (%s)" % (b["n_text"], _list_wells(b["text_wells"])))
        lines.append("Block %s on %s at %s: %d wells, %d with a value%s." % (
            b["block"], b["plate"], b["range"], b["n_wells"], b["n_values"], ", missing: " + "; ".join(miss) if miss else ", none missing"))
    by_order = [b["block"] for b in blocks if b["letters_from_order"]]
    if by_order:
        lines.append("The file has no row letters for %d blocks (%s). The row letters came from the row order, A for the first row." % (len(by_order), "; ".join(by_order)))
    if skipped:
        lines.append("Skipped text-only grids: %s." % "; ".join(skipped))
    if mapping:
        lines.append("Plate map: %s." % mapping[2])
        for b in blocks:
            if b["n_map_without_value"]:
                lines.append("Block %s, mapped wells without a value: %d (%s)." % (b["block"], b["n_map_without_value"], _list_wells(b["map_without_value"])))
        if not any(b["n_map_without_value"] for b in blocks):
            lines.append("Every mapped well has a value.")
    lines.append("A missing value stays empty in the file. It is not zero. The column flag shows the text of the cell.")
    lines.append("Saved %s (%d rows, columns %s)." % (out, len(frame), ", ".join(map(str, frame.columns))))
    route = "In Excel: open %s. " % os.path.basename(path)
    route += "Select each block with its row letters and column numbers (%s). " % "; ".join("%s on %s" % (b["range"], b["plate"]) for b in blocks)
    route += "Use Data > From Table/Range and make it a table with headers. In Power Query select the row letter column and use Transform > Unpivot Other Columns. "
    route += "Rename Attribute to col and Value to value. Add a column well with the formula [Row] & [col]. "
    route += "Text cells such as OVRFLW stay text. Filter them out. Do not turn them into zero. "
    route += "Append the blocks to one query and add a column block with the label of each block."
    if mapping:
        route += " Unpivot the plate map grid the same way, or load its long table. Then use Home > Merge Queries on the column well (left outer join)."
    return {
        "ok": True,
        "summary": " ".join(lines),
        "metrics": {
            "n_blocks": len(blocks), "n_plates": len({b["plate"] for b in blocks}), "n_wells": sum(b["n_wells"] for b in blocks), "n_values": sum(b["n_values"] for b in blocks),
            "n_missing": sum(b["n_wells"] - b["n_values"] for b in blocks), "n_empty": sum(b["n_empty"] for b in blocks), "n_overflow": sum(b["n_overflow"] for b in blocks),
            "n_text": sum(b["n_text"] for b in blocks), "n_map_without_value": sum(b["n_map_without_value"] for b in blocks), "n_map_columns": n_map_cols,
        },
        "path": out, "blocks": blocks, "skipped": skipped, "columns": _col_types(frame),
        "table": _preview(frame, out),
        "outputs": [{"path": out, "kind": "table", "name": os.path.basename(out)}],
        "manualRoute": route,
        "fidelity": "close",
    }
