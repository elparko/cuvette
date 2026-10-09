import os
from datetime import date

from openpyxl import Workbook

here = os.path.dirname(os.path.abspath(__file__))
wb = Workbook()
ws = wb.active
ws.title = "Counts"
ws["A1"] = "RNAscope counts, batch 4 (synthetic test data)"
ws["D3"] = "Control"
ws.merge_cells("D3:E3")
ws["F3"] = "Knockdown"
ws.merge_cells("F3:G3")
for col, text in zip("CDEFG", ["Sample", "Red dots", "Total cells", "Red dots", "Total cells"]):
    ws[col + "5"] = text
control = [10, 12, 14, 16, 18, 20]
knockdown = [5, 6, 4, 5, 7, 3]
for i in range(6):
    r = 6 + i
    ws["C%d" % r] = "Slice %d" % (i + 1)
    ws["D%d" % r], ws["E%d" % r] = control[i], 100
    ws["F%d" % r], ws["G%d" % r] = knockdown[i], 50
ws["C12"], ws["D12"] = "Slice 7", "lost"
ws["C13"], ws["D13"], ws["E13"], ws["F13"], ws["G13"] = "Average", 15, 100, 5, 50
ws["C15"] = "Note: slice 4 knockdown was imaged twice, the second image is used"
ws["I3"] = "Imaging"
ws["I5"], ws["J5"] = "Date", "Slide"
for i, (d, s) in enumerate([(date(2026, 9, 1), "S-101"), (date(2026, 9, 2), "S-102"), (date(2026, 9, 3), "S-103")]):
    ws["I%d" % (6 + i)], ws["J%d" % (6 + i)] = d, s
ws["L10"] = "ask PI about slide 3"
wb.create_sheet("Plate map")["A1"] = "empty plate map"
wb.save(os.path.join(here, "messy_lab.xlsx"))

with open(os.path.join(here, "slices.csv"), "w") as f:
    f.write("Sample,Section depth (um)\n")
    for i in range(5):
        f.write("Slice %d,%d\n" % (i + 1, 40 + 10 * i))


def letters(n):
    return [chr(65 + i) for i in range(n)]


def gen5_value(i, j, scale=1):
    return round((0.1 * (j + 1) + 0.01 * i) * scale, 4)


def gen5_block(ws, top, label, scale, special):
    ws.cell(top, 1, label)
    for j in range(12):
        ws.cell(top + 1, 2 + j, j + 1)
    for i, letter in enumerate(letters(8)):
        ws.cell(top + 2 + i, 1, letter)
        for j in range(12):
            ws.cell(top + 2 + i, 2 + j, special.get((i, j), gen5_value(i, j, scale)))


wb = Workbook()
ws = wb.active
ws.title = "Gen5"
for r, (a, b) in enumerate([("Software Version", "3.11"), ("Experiment File Path:", "C:\\synthetic\\elisa.xpt"), ("Plate Number", "Plate 1"), ("Date", "2026-09-01"), ("Reader Type:", "Synthetic reader")], 1):
    ws.cell(r, 1, a)
    ws.cell(r, 2, b)
gen5_block(ws, 8, "Read 1:450", 1, {(7, 11): "OVRFLW"})
gen5_block(ws, 20, "Read 2:630", 0.1, {(6, 2): None, (3, 5): "*"})
pm = wb.create_sheet("Plate map")
pm["A1"] = "Plate map (synthetic)"
for j in range(12):
    pm.cell(3, 2 + j, j + 1)
standards = ["STD1", "STD 2", "Std3", "STD4", "STD5", "STD6", "BLK", "BLK"]
for i, letter in enumerate(letters(8)):
    pm.cell(4 + i, 1, letter)
    for j in range(12):
        pm.cell(4 + i, 2 + j, standards[i] if j < 2 else "S%02d" % (j - 1) if j < 10 else "NC")
wb.save(os.path.join(here, "plate_gen5.xlsx"))

with open(os.path.join(here, "plate_softmax.csv"), "w") as f:
    f.write("Plate:,Plate#1,1.3,PlateFormat,Endpoint,Absorbance,Raw,TRUE,1,,,450,,,1,12,96,1,8,,,,\n")
    f.write(",,Temperature(C),%s,\n" % ",".join(str(j + 1) for j in range(12)))
    for i, letter in enumerate(letters(8)):
        f.write(",%s,%s,%s,\n" % (letter, "24.3" if i == 0 else "", ",".join(str(round(0.05 * (i + 1) + 0.002 * j, 4)) for j in range(12))))
    f.write("\n~End\n")

with open(os.path.join(here, "plate_softmax_map.csv"), "w") as f:
    f.write("Well,Sample,Conc_uM\n")
    for i, letter in enumerate("AB"):
        for j in range(3):
            f.write("%s%02d,Std %d,%d\n" % (letter, j + 1, i * 3 + j + 1, 10 * (i * 3 + j + 1)))
    f.write("J1,Extra,0\n")

with open(os.path.join(here, "plate_384.csv"), "w") as f:
    f.write("Tecan i-control,1.10\nDevice: synthetic\n\nMode,Absorbance\n\nLabel: Label1\n")
    f.write("<>,%s\n" % ",".join(str(j + 1) for j in range(24)))
    for i, letter in enumerate(letters(16)):
        f.write("%s,%s\n" % (letter, ",".join(str(round((i * 24 + j + 1) / 1000, 3)) for j in range(24))))
    f.write("\nEnd Time:,2026-09-01\n")
