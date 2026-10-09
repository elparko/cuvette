import json
import sys

from PIL import Image


def to_png(src, dst):
    with Image.open(src) as im:
        im.seek(0)
        if im.mode in ("I;16", "I;16B", "I;16L", "I", "F"):
            import numpy as np

            a = np.asarray(im).astype("float64")
            lo, hi = float(a.min()), float(a.max())
            a = (a - lo) / (hi - lo or 1) * 255
            out = Image.fromarray(a.clip(0, 255).astype("uint8"))
        elif im.mode in ("RGB", "RGBA", "L", "LA", "P"):
            out = im.copy()
        else:
            out = im.convert("RGB")
        out.save(dst, "PNG")


def pptx(spec):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Mm, Pt

    prs = Presentation()
    prs.slide_width = Mm(spec["size"][0])
    prs.slide_height = Mm(spec["size"][1])
    blank = prs.slide_layouts[6]
    for page in spec["pages"]:
        slide = prs.slides.add_slide(blank)
        for c in page["images"]:
            x, y, w, h = c["rect"]
            with Image.open(c["path"]) as im:
                iw, ih = im.size
            s = min(w / iw, h / ih)
            dw, dh = iw * s, ih * s
            slide.shapes.add_picture(c["path"], Mm(x + (w - dw) / 2), Mm(y + h - dh), Mm(dw), Mm(dh))
        for t in page["texts"]:
            x, y, w, h = t["rect"]
            tf = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h)).text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            for i, p in enumerate(t["paras"]):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                if p.get("space"):
                    para.space_before = Pt(p["space"])
                run = para.add_run()
                run.text = p["text"]
                run.font.size = Pt(p["pt"])
                run.font.bold = bool(p.get("bold"))
                run.font.name = "Arial"
                run.font.color.rgb = RGBColor(0, 0, 0)
        if page.get("notes"):
            slide.notes_slide.notes_text_frame.text = page["notes"]
    prs.save(spec["out"])


def main():
    with open(sys.argv[1], encoding="utf-8") as f:
        spec = json.load(f)
    for c in spec.get("convert", []):
        try:
            to_png(c["src"], c["dst"])
        except Exception as e:
            print(f"could not convert {c['src']}: {e}", file=sys.stderr)
    if spec.get("pages"):
        pptx(spec)


main()
